"""Aggregate the raw per-utterance scores from detect_eval.py.

Reads every ``<Method>_<Dataset>.npz`` under --in and produces

  * an Excel workbook: one sheet per evaluation corpus with bit accuracy
    (mean, std, 95% CI), AUC (bootstrap CI), and TPR/FPR at a calibrated
    threshold, plus a Summary sheet averaged over the four corpora;
  * the LaTeX body for tab:distortion_comparison.

Detection statistic
-------------------
The number of payload bits recovered. Positives are watermarked clips scored
against the embedded message; negatives are the same clips *unwatermarked*,
scored against an independent random message. Under H0 the statistic is
Binomial(n_bits, 0.5), so the negatives are an empirical null.

Threshold
---------
One threshold per method, not per cell. A deployed detector has a single
operating point, and a per-cell threshold would be fitted to ~2300 negatives
whose 1% quantile rests on ~23 samples. Pooling every negative a method
produced (4 corpora x 10 channels, ~86k draws) resolves that quantile far
better, and unwatermarked audio is close to Binomial(n, 0.5) whatever channel
it passed through. The threshold is the smallest integer whose pooled FPR is
at or below the target, so the realised FPR lands under 1% rather than on it:
the statistic is discrete and cannot hit an arbitrary rate exactly.

    python detect_report.py --in results/detect --out results/evals
"""

import argparse
import glob
import os
import re

import numpy as np
from sklearn.metrics import roc_auc_score

# Display name -> key used by detect_eval, in table order.
ROWS = [
    ("No Distortion", "none"),
    ("Resample (8k)", "resample_8k"),
    ("Gaussian Noise", "gaussian_noise_20"),
    ("Median Filter", "median_filter"),
    ("Low-pass 4k", "low_pass_4k"),
    ("High-pass 500", "high_pass_500"),
    ("Reencode", "reencode"),
    ("Compression", "compression"),
    ("Noise Suppression", "noise_suppression"),
    ("Phone Call", "phone_call"),
]
METHODS = ["RT-SW", "AudioSeal", "WavMark", "Timbre", "SilentCipher"]
# payload each detector recovers; the imperceptibility table needs these
# even before the matching detection run exists
PAYLOAD = {"RT-SW": 10, "AudioSeal": 16, "WavMark": 16, "Timbre": 10,
           "SilentCipher": 40}
DATASETS = ["LibriSpeech-dev", "LJSpeech", "clone_xspeaker", "resynth_hifigan"]
TARGET_FPR = 0.01
Z = 1.959963984540054  # two-sided 95%


# ------------------------------------------------------------------ statistics


def wilson(k, n, z=Z):
    """Wilson score interval. Normal approximation collapses at p near 0 or 1,
    which is exactly where TPR and FPR sit here."""
    if n == 0:
        return (float("nan"), float("nan"))
    p = k / n
    d = 1 + z * z / n
    c = (p + z * z / (2 * n)) / d
    h = z * np.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / d
    return (max(0.0, c - h), min(1.0, c + h))


def auc_ci(pos, neg, n_boot=2000, seed=0):
    """AUC with a paired cluster bootstrap CI.

    Each clip contributes one positive and one negative at the same index, so
    the two classes are correlated: a hard utterance drags both down together.
    Resampling the classes independently would throw that pairing away and
    misstate the interval, so the bootstrap resamples *clips*, carrying each
    clip's positive and negative along together.
    """
    pos = np.asarray(pos, float)
    neg = np.asarray(neg, float)
    if len(pos) == 0 or len(neg) == 0:
        return float("nan"), float("nan"), float("nan")
    assert len(pos) == len(neg), (len(pos), len(neg))
    n = len(pos)
    y = np.r_[np.ones(n), np.zeros(n)]
    if len(np.unique(np.r_[pos, neg])) == 1:  # degenerate: all scores identical
        return 0.5, 0.5, 0.5
    point = roc_auc_score(y, np.r_[pos, neg])
    rng = np.random.default_rng(seed)
    boot = np.empty(n_boot)
    for b in range(n_boot):
        idx = rng.integers(0, n, n)
        p, q = pos[idx], neg[idx]
        boot[b] = 0.5 if len(np.unique(np.r_[p, q])) == 1 \
            else roc_auc_score(y, np.r_[p, q])
    lo, hi = np.percentile(boot, [2.5, 97.5])
    return float(point), float(lo), float(hi)


def calibrate(neg, n_bits, target=TARGET_FPR):
    """Smallest integer threshold whose FPR on `neg` is <= target."""
    for t in range(0, n_bits + 1):
        if float(np.mean(neg >= t)) <= target:
            return t
    return n_bits + 1


def calibrate_randomized(neg, n_bits, target=TARGET_FPR):
    """Randomised Neyman-Pearson threshold hitting the target FPR exactly.

    The statistic is an integer bit count, so the attainable FPRs are a ladder
    whose spacing is set by the payload size. At 10 bits the only rungs near 1%
    are 1.07% (t=9) and 0.098% (t=10), so a deterministic threshold holds a
    10-bit method to a rate 10x stricter than a 40-bit one and its TPR is not
    comparable. The standard remedy rejects outright above t and with
    probability gamma exactly at t-1, which lands on the target for every
    payload size and makes TPRs comparable across methods.
    """
    t = calibrate(neg, n_bits, target)
    above = float(np.mean(neg >= t))
    at = float(np.mean(neg == t - 1)) if t >= 1 else 0.0
    gamma = 0.0 if at <= 0 else min(1.0, max(0.0, (target - above) / at))
    return t, gamma


def apply_randomized(scores, t, gamma):
    """Expected rejection rate of the randomised rule (exact, not simulated)."""
    scores = np.asarray(scores, float)
    if len(scores) == 0:
        return float("nan")
    return float(np.mean(scores >= t) + gamma * np.mean(scores == t - 1))


def cell(pos, neg, n_bits, thr, rand=None, seed=0):
    pos = np.asarray(pos, float)
    neg = np.asarray(neg, float)
    n = len(pos)
    acc = pos / n_bits
    mean = float(acc.mean()) if n else float("nan")
    std = float(acc.std(ddof=1)) if n > 1 else float("nan")
    sem = std / np.sqrt(n) if n > 1 else float("nan")
    a, alo, ahi = auc_ci(pos, neg, seed=seed)
    ktp, ntp = int((pos >= thr).sum()), n
    kfp, nfp = int((neg >= thr).sum()), len(neg)
    tlo, thi = wilson(ktp, ntp)
    flo, fhi = wilson(kfp, nfp)
    return dict(
        n=n, acc=mean, acc_std=std, acc_sem=sem,
        acc_lo=mean - Z * sem if n > 1 else float("nan"),
        acc_hi=mean + Z * sem if n > 1 else float("nan"),
        auc=a, auc_lo=alo, auc_hi=ahi,
        tpr=ktp / ntp if ntp else float("nan"), tpr_lo=tlo, tpr_hi=thi,
        fpr=kfp / nfp if nfp else float("nan"), fpr_lo=flo, fpr_hi=fhi,
        thr=thr,
        # exactly-1%-FPR operating point, comparable across payload sizes
        tpr_x=apply_randomized(pos, *rand) if rand else float("nan"),
        fpr_x=apply_randomized(neg, *rand) if rand else float("nan"),
    )


# ------------------------------------------------------------------ loading


def load(indir):
    """{method: {dataset: {'pos'/'neg': {row: array}, 'n_bits': int}}}

    A corpus may be split across several files. One process per corpus leaves
    most of the machine idle once only a few corpora remain -- five workers at
    98% of one core each with nineteen cores free -- so a corpus can be sharded
    over disjoint clip ranges. Shards are named
    ``<Method>_<Corpus>__shardN.npz``; because their clips are disjoint,
    concatenating them reproduces the whole-corpus result exactly.
    """
    parts = {}
    for f in sorted(glob.glob(os.path.join(indir, "*.npz"))):
        base = os.path.basename(f)[:-4]
        if base.startswith("_"):
            continue
        base = re.sub(r"__shard\d+$", "", base)
        m = re.match(r"^(.*)_(%s)$" % "|".join(map(re.escape, DATASETS)), base)
        if not m:
            print(f"  skipping unrecognised {base}")
            continue
        ds = m.group(2)
        d = np.load(f, allow_pickle=True)
        meth = str(d["method"])
        lens = [len(d[f"{w}_{k}"]) for _, k in ROWS
                for w in ("pos", "neg") if f"{w}_{k}" in d]
        k_min = min(lens) if lens else 0
        if lens and max(lens) != k_min:
            print(f"  {meth}/{ds}: ragged ({k_min}..{max(lens)}), truncating")
        complete = bool(d["complete"]) if "complete" in d.files else True
        parts.setdefault((meth, ds), []).append((
            int(d["n_bits"]), complete,
            {k: d[f"pos_{k}"][:k_min] for _, k in ROWS if f"pos_{k}" in d},
            {k: d[f"neg_{k}"][:k_min] for _, k in ROWS if f"neg_{k}" in d}))

    data = {}
    for (meth, ds), chunks in parts.items():
        n_bits = chunks[0][0]
        pos = {k: np.concatenate([c[2][k] for c in chunks if k in c[2]])
               for _, k in ROWS if any(k in c[2] for c in chunks)}
        neg = {k: np.concatenate([c[3][k] for c in chunks if k in c[3]])
               for _, k in ROWS if any(k in c[3] for c in chunks)}
        used = len(next(iter(pos.values()))) if pos else 0
        if len(chunks) > 1:
            print(f"  {meth}/{ds}: merged {len(chunks)} shards -> {used} clips")
        elif not chunks[0][1]:
            print(f"  {meth}/{ds}: partial, {used} clips so far")
        data.setdefault(meth, {})[ds] = dict(n_bits=n_bits, used=used,
                                             pos=pos, neg=neg)
    return data


HDR = ["distortion", "n", "bit_acc", "acc_std", "acc_sem", "acc_ci_lo",
       "acc_ci_hi", "auc", "auc_ci_lo", "auc_ci_hi", "tpr", "tpr_ci_lo",
       "tpr_ci_hi", "fpr", "fpr_ci_lo", "fpr_ci_hi", "threshold_bits",
       "tpr_at_exactly_1pct_fpr", "fpr_randomized"]

# HDR carries display names; KEYS are the matching cell() dict keys, in order.
KEYS = ["acc", "acc_std", "acc_sem", "acc_lo", "acc_hi",
        "auc", "auc_lo", "auc_hi", "tpr", "tpr_lo", "tpr_hi",
        "fpr", "fpr_lo", "fpr_hi", "thr", "tpr_x", "fpr_x"]
assert len(KEYS) == len(HDR) - 2


def _sheet(ws, title):
    ws.title = title
    return ws


def write_xlsx(path, per, summary, thresholds, order, imp=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    bold = Font(bold=True)

    ws = wb.active
    ws.title = "Summary"
    ws.append(["Averaged over the four evaluation corpora (unweighted mean of "
               "per-corpus values; CI columns are pooled over all clips)"])
    ws["A1"].font = bold
    ws.append([])
    for meth in order:
        ws.append([meth])
        ws.cell(ws.max_row, 1).font = bold
        ws.append(HDR)
        for c in range(1, len(HDR) + 1):
            ws.cell(ws.max_row, c).font = bold
        for label, key in ROWS:
            r = summary[meth][key]
            ws.append([label, r["n"]] + [round(r[k], 6) for k in KEYS])
        ws.append([])

    for ds in order_datasets(per):
        ws = wb.create_sheet(ds[:31])
        ws.append([f"{ds}"])
        ws["A1"].font = bold
        ws.append([])
        for meth in order:
            if ds not in per.get(meth, {}):
                continue
            ws.append([meth])
            ws.cell(ws.max_row, 1).font = bold
            ws.append(HDR)
            for c in range(1, len(HDR) + 1):
                ws.cell(ws.max_row, c).font = bold
            for label, key in ROWS:
                r = per[meth][ds].get(key)
                if r is None:
                    continue
                ws.append([label, r["n"]] + [round(r[k], 6) for k in KEYS])
            ws.append([])

    ws = wb.create_sheet("Thresholds")
    ws.append(["method", "payload_bits", "threshold_bits", "pooled_negatives",
               "realised_pooled_fpr", "target_fpr", "randomized_t",
               "randomized_gamma"])
    for c in range(1, 9):
        ws.cell(1, c).font = bold
    for meth in order:
        t = thresholds[meth]
        ws.append([meth, t["n_bits"], t["thr"], t["n_neg"],
                   round(t["fpr"], 6), TARGET_FPR, t["rand"][0],
                   round(t["rand"][1], 6)])
    ws.append([])
    ws.append(["Deterministic threshold: smallest integer bit count whose "
               "pooled FPR is <= the target. Because the statistic is discrete "
               "the realised rate lands under the target, by a margin that "
               "depends on payload size."])
    ws.append(["Randomized rule: reject when score >= randomized_t, or when "
               "score == randomized_t - 1 with probability randomized_gamma. "
               "Hits the target FPR exactly for every payload size, so TPRs "
               "are comparable across methods."])

    if imp:
        ws = wb.create_sheet("Imperceptibility")
        ws.append(["Measured on the clean/watermarked pair before any channel: "
                   "no distortion, no decoding, no threshold."])
        ws["A1"].font = bold
        ws.append([])
        hdr = ["method", "dataset", "n", "snr_mean_db", "snr_std",
               "snr_ci_lo", "snr_ci_hi", "pesq_mean", "pesq_std",
               "pesq_ci_lo", "pesq_ci_hi"]
        ws.append(hdr)
        for c in range(1, len(hdr) + 1):
            ws.cell(ws.max_row, c).font = bold
        for meth in order:
            if meth not in imp:
                continue
            rows = []
            for ds in order_datasets(per) or list(imp[meth]):
                if ds not in imp[meth]:
                    continue
                st = imp_stats(imp[meth][ds])
                rows.append(st)
                ws.append([meth, ds, st["snr"]["n"]]
                          + [round(st["snr"][k], 4)
                             for k in ("mean", "std", "lo", "hi")]
                          + [round(st["pesq"][k], 4)
                             for k in ("mean", "std", "lo", "hi")])
            if rows:
                ws.append([meth, "MEAN OF CORPORA", "",
                           round(float(np.mean([r["snr"]["mean"] for r in rows])), 4),
                           "", "", "",
                           round(float(np.mean([r["pesq"]["mean"] for r in rows])), 4)])
                ws.cell(ws.max_row, 1).font = bold

    ws = wb.create_sheet("Raw")
    ws.append(["method", "dataset"] + HDR)
    for c in range(1, len(HDR) + 3):
        ws.cell(1, c).font = bold
    for meth in order:
        for ds in order_datasets(per):
            for label, key in ROWS:
                r = per.get(meth, {}).get(ds, {}).get(key)
                if r is None:
                    continue
                ws.append([meth, ds, label, r["n"]]
                          + [round(r[k], 6) for k in KEYS])

    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb.save(path)


def load_imp(indir):
    """{method: {dataset: {'snr': array, 'pesq': array}}}"""
    out = {}
    for f in sorted(glob.glob(os.path.join(indir, "*.npz"))):
        base = os.path.basename(f)[:-4]
        if base.startswith("_"):
            continue
        m = re.match(r"^(.*)_(%s)$" % "|".join(map(re.escape, DATASETS)), base)
        if not m:
            continue
        d = np.load(f, allow_pickle=True)
        out.setdefault(str(d["method"]), {})[m.group(2)] = dict(
            snr=np.asarray(d["snr"], float), pesq=np.asarray(d["pesq"], float))
    return out


def imp_stats(v):
    """Mean, std and 95% CI of SNR and PESQ, ignoring clips PESQ rejected."""
    r = {}
    for k in ("snr", "pesq"):
        a = np.asarray(v[k], float)
        a = a[np.isfinite(a)]
        n = len(a)
        mean = float(a.mean()) if n else float("nan")
        std = float(a.std(ddof=1)) if n > 1 else float("nan")
        sem = std / np.sqrt(n) if n > 1 else float("nan")
        r[k] = dict(n=n, mean=mean, std=std,
                    lo=mean - Z * sem if n > 1 else float("nan"),
                    hi=mean + Z * sem if n > 1 else float("nan"))
    return r


IMP_TEX = r"""\begin{table}[t]
\centering
\caption{Imperceptibility of the embedded watermark, averaged over the four
evaluation corpora. SNR is speech power over watermark power; PESQ is ITU-T
P.862 wideband. Payload is the message size each detector recovers.}
\label{tab:imperceptibility}
\begin{tabular}{lccc}
\toprule
\textbf{Method} & \textbf{Payload (bits)} & \textbf{SNR (dB)} & \textbf{PESQ} \\
\midrule
"""


def write_imp_tex(path, imp, thresholds):
    lines = []
    for meth in METHODS:
        if meth not in imp:
            continue
        dss = list(imp[meth])
        st = [imp_stats(imp[meth][ds]) for ds in dss]
        snr = np.mean([x["snr"]["mean"] for x in st])
        pq = np.mean([x["pesq"]["mean"] for x in st])
        nb = thresholds.get(meth, {}).get("n_bits") or PAYLOAD.get(meth, "--")
        lines.append(f"{meth} & {nb} & {snr:.2f} & {pq:.3f} " + r"\\")
    open(path, "w").write(IMP_TEX + "\n".join(lines) +
                          "\n\\bottomrule\n\\end{tabular}\n\\end{table}\n")


def order_datasets(per):
    seen = [d for d in DATASETS if any(d in v for v in per.values())]
    return seen


# ------------------------------------------------------------------ LaTeX

TEX_HEAD = r"""\begin{table*}[t]
\centering
\caption{
Detection results under different audio distortions.
Acc. denotes detection accuracy, TPR/FPR denotes the true-positive
and false-positive rates, and AUC denotes the area under the ROC curve.
}
\label{tab:distortion_comparison}

\resizebox{\textwidth}{!}{
\begin{tabular}{l ccc ccc ccc ccc ccc}
\toprule
& \multicolumn{3}{c}{\textbf{RT-SW}}
& \multicolumn{3}{c}{\textbf{AudioSeal}}
& \multicolumn{3}{c}{\textbf{WavMark}}
& \multicolumn{3}{c}{\textbf{Timbre}}
& \multicolumn{3}{c}{\textbf{SilentCipher}} \\

\cmidrule(lr){2-4}
\cmidrule(lr){5-7}
\cmidrule(lr){8-10}
\cmidrule(lr){11-13}
\cmidrule(lr){14-16}

\textbf{Distortion}
& \textbf{Acc.} & \textbf{TPR/FPR} & \textbf{AUC}
& \textbf{Acc.} & \textbf{TPR/FPR} & \textbf{AUC}
& \textbf{Acc.} & \textbf{TPR/FPR} & \textbf{AUC}
& \textbf{Acc.} & \textbf{TPR/FPR} & \textbf{AUC}
& \textbf{Acc.} & \textbf{TPR/FPR} & \textbf{AUC} \\
\midrule
"""

TEX_TAIL = r"""
\bottomrule
\end{tabular}
}
\end{table*}
"""


def fmt(v, nd=3):
    return "--" if v != v else f"{v:.{nd}f}"


def write_tex(path, summary, mode="deterministic"):
    tk, fk = ("tpr", "fpr") if mode == "deterministic" else ("tpr_x", "fpr_x")
    lines = []
    for label, key in ROWS:
        parts = [label]
        for meth in METHODS:
            r = summary.get(meth, {}).get(key)
            if r is None:
                parts.append("& -- & -- & --")
            else:
                parts.append(f"& {fmt(r['acc'])} & {fmt(r[tk])}/"
                             f"{fmt(r[fk], 4)} & {fmt(r['auc'])}")
        lines.append("\n".join(parts) + r" \\")
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    open(path, "w").write(TEX_HEAD + "\n\n".join(lines) + TEX_TAIL)


# ------------------------------------------------------------------ main


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--in", dest="indir", default="results/detect")
    ap.add_argument("--imp", dest="impdir", default="results/imp",
                    help="directory of imperceptibility.py outputs; '' to skip")
    ap.add_argument("--out", default="results/evals")
    ap.add_argument("--n_boot", type=int, default=2000)
    ap.add_argument("--tex_tpr", choices=["deterministic", "exact"],
                    default="deterministic",
                    help="which operating point the LaTeX TPR/FPR column uses")
    args = ap.parse_args()

    data = load(args.indir)
    imp = load_imp(args.impdir) if args.impdir and os.path.isdir(args.impdir) else {}
    if not data:
        raise SystemExit(f"no .npz under {args.indir}")

    order = [m for m in METHODS if m in data] + \
            [m for m in sorted(data) if m not in METHODS]

    # One threshold per method, from every negative it produced.
    thresholds = {}
    for meth, byds in data.items():
        n_bits = next(iter(byds.values()))["n_bits"]
        pooled = np.concatenate([v["neg"][k] for v in byds.values()
                                 for _, k in ROWS if k in v["neg"]])
        thr = calibrate(pooled, n_bits)
        rand = calibrate_randomized(pooled, n_bits)
        thresholds[meth] = dict(n_bits=n_bits, thr=thr, n_neg=len(pooled),
                                fpr=float(np.mean(pooled >= thr)), rand=rand)
        print(f"{meth:<13} payload {n_bits:>2} bits, threshold >= {thr} bits, "
              f"pooled FPR {thresholds[meth]['fpr']:.4f} on {len(pooled)} "
              f"negatives; randomised rule >= {rand[0]} or == {rand[0]-1} "
              f"w.p. {rand[1]:.4f} -> exactly {TARGET_FPR:.1%}")

    per, summary = {}, {}
    for meth, byds in data.items():
        n_bits = thresholds[meth]["n_bits"]
        thr = thresholds[meth]["thr"]
        rand = thresholds[meth]["rand"]
        per[meth] = {}
        for ds, v in byds.items():
            per[meth][ds] = {
                k: cell(v["pos"][k], v["neg"][k], n_bits, thr, rand)
                for _, k in ROWS if k in v["pos"]
            }
        # Averaged row: pooled clips give the CI, mean-of-corpora the headline.
        summary[meth] = {}
        for _, k in ROWS:
            dss = [ds for ds in byds if k in byds[ds]["pos"]]
            if not dss:
                continue
            pos = np.concatenate([byds[ds]["pos"][k] for ds in dss])
            neg = np.concatenate([byds[ds]["neg"][k] for ds in dss])
            r = cell(pos, neg, n_bits, thr, rand)
            for f in ("acc", "auc", "tpr", "fpr", "tpr_x", "fpr_x"):
                r[f] = float(np.mean([per[meth][ds][k][f] for ds in dss]))
            summary[meth][k] = r

    xlsx = os.path.join(args.out, "detection_metrics.xlsx")
    tex = os.path.join(args.out, "distortion_comparison.tex")
    write_xlsx(xlsx, per, summary, thresholds, order, imp)
    # Both operating points, always. The deterministic threshold is what a
    # deployed detector does; the randomised one is the only comparable
    # TPR across payload sizes. Writing both saves re-running to switch.
    write_tex(tex, summary, "deterministic")
    tex_x = os.path.join(args.out, "distortion_comparison_exact1pct.tex")
    write_tex(tex_x, summary, "exact")
    if imp:
        write_imp_tex(os.path.join(args.out, "imperceptibility.tex"), imp,
                      thresholds)

    print(f"\nwrote {xlsx}\nwrote {tex}  (deterministic threshold)"
          f"\nwrote {tex_x}  (randomised, exactly 1% FPR)\n")
    hdr = f"{'distortion':<20}" + "".join(f"{m:>26}" for m in order)
    print(hdr)
    print("-" * len(hdr))
    for label, key in ROWS:
        row = f"{label:<20}"
        for meth in order:
            r = summary.get(meth, {}).get(key)
            row += "  " + (f"{'--':>24}" if r is None else
                           f"{r['acc']:.3f} {r['tpr']:.3f}/{r['fpr']:.3f} "
                           f"{r['auc']:.3f}".rjust(24))
        print(row)
    missing = [(m, d) for m in order for d in DATASETS if d not in per.get(m, {})]
    if missing:
        print("\nstill missing: " + ", ".join(f"{m}/{d}" for m, d in missing))


if __name__ == "__main__":
    main()
